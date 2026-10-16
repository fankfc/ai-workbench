"""Real file rendering, font coverage, byte/text and geometry verification.

DOCX without a bundled office renderer remains blocked even if the package is
structurally valid. A PDF drawn from the same model is not a DOCX page preview.
"""
from dataclasses import dataclass
import io
import os
from pathlib import Path
import re
import tempfile
import zipfile
import subprocess
import xml.etree.ElementTree as ET

from .model import DocumentModel, digest


@dataclass(frozen=True)
class RenderedDocument:
    format: str
    language: str
    data: bytes
    qa: dict
    page_images: tuple[Path, ...] = ()
    preview_pdf: bytes | None = None


def _visible(text):
    from markdown_it import MarkdownIt
    values=[]
    for token in MarkdownIt('commonmark').parseInline(text):
        for child in token.children or []:
            if child.type in ('text','code_inline'): values.append(child.content)
            elif child.type in ('softbreak','hardbreak'): values.append('\n')
    return ''.join(values)


def _block_text(b): return b.content if b.kind=='code' else _visible(b.content)


def _compact(text):
    return ''.join(text.split())


def _pdf_rich(text):
    from xml.sax.saxutils import escape
    values=[]; offset=0
    for match in re.finditer(r'\[([^]]+)\]\((https?://[^)\s]+|mailto:[^)\s]+)\)',text):
        values.append(escape(_visible(text[offset:match.start()])))
        values.append('<link href="'+escape(match[2],{'"':'&quot;'})+'" color="#225a91">'+escape(match[1])+'</link>')
        offset=match.end()
    values.append(escape(_visible(text[offset:])))
    return ''.join(values).replace('\n','<br/>')


class WindowsWordRenderer:
    """Hidden Word-compatible COM export of controlled OOXML, serialized locally."""
    def __call__(self, data, root):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if any('vba' in n.lower() for n in archive.namelist()): raise ValueError('DOCX_MACROS_FORBIDDEN')
            for name in archive.namelist():
                if not name.endswith('.rels'): continue
                for relation in ET.fromstring(archive.read(name)):
                    if relation.get('TargetMode')=='External' and (not relation.get('Type','').endswith('/hyperlink')
                            or re.match(r'^(https?://|mailto:)',relation.get('Target','')) is None):
                        raise ValueError('DOCX_EXTERNAL_RELATIONSHIP_FORBIDDEN')
        directory = Path(tempfile.mkdtemp(prefix='word-render-', dir=root))
        source, output, script = directory/'input.docx', directory/'input.pdf', directory/'render.ps1'
        source.write_bytes(data)
        script.write_text('''param([string]$SourceFile, [string]$OutputFile)
$ErrorActionPreference = "Stop"
$wordApp = $null
$wordDoc = $null
$saveOption = 0
$exportCompleted = $false
$documentClosed = $false
$renderMutex = New-Object System.Threading.Mutex($false, 'Local\\AIWorkbenchDocxRender')
$mutexOwned = $false
try {
    try { $mutexOwned = $renderMutex.WaitOne(30000) }
    catch [System.Threading.AbandonedMutexException] { $mutexOwned = $true }
    if (-not $mutexOwned) { throw 'DOCX_RENDERER_BUSY' }
    $wordApp = New-Object -ComObject Word.Application
    $wordApp.Visible = $false
    $wordApp.DisplayAlerts = 0
    $wordApp.AutomationSecurity = 3
    $wordDoc = $wordApp.Documents.Open($SourceFile, $false, $true, $false)
    $wordDoc.ExportAsFixedFormat($OutputFile, 17)
    $exportCompleted = $true
    Write-Output $wordApp.Version
} finally {
    if ($null -ne $wordDoc) {
        $wordDoc.Close([ref]$saveOption)
        $documentClosed = $true
        [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($wordDoc)
    }
    if ($null -ne $wordApp) {
        try { $wordApp.Quit([ref]$saveOption) }
        catch {
            # Compatible hosts may disconnect after closing the document.
            # Accept only those RPC errors after a successful export/close;
            # output bytes still undergo independent content/layout checks.
            $rpcDisconnected = $false
            $cleanupException = $_.Exception
            while ($null -ne $cleanupException) {
                if ($cleanupException.HResult -in @(-2147023174, -2147023170)) { $rpcDisconnected = $true }
                $cleanupException = $cleanupException.InnerException
            }
            if (-not ($exportCompleted -and $documentClosed -and $rpcDisconnected)) { throw }
        }
        [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($wordApp)
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
    if ($mutexOwned) { $renderMutex.ReleaseMutex() }
    $renderMutex.Dispose()
}
''', encoding='utf-8-sig')
        command = [str(Path(os.environ.get('WINDIR','C:/Windows'))/'System32/WindowsPowerShell/v1.0/powershell.exe'),
            '-NoProfile','-NonInteractive','-ExecutionPolicy','Bypass','-File',str(script),str(source),str(output)]
        try:
            result = subprocess.run(command, capture_output=True, text=True, timeout=45, check=False)
        except subprocess.TimeoutExpired as error:
            raise ValueError('DOCX_WORD_RENDER_TIMEOUT') from error
        (directory/'renderer.log').write_text(result.stdout+'\n'+result.stderr, encoding='utf-8')
        if result.returncode or not output.is_file(): raise ValueError('DOCX_WORD_RENDER_FAILED')
        return output.read_bytes()


def _native_docx_renderer():
    if os.name != 'nt': return None
    import winreg
    try:
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r'Word.Application\CLSID'):
            return WindowsWordRenderer()
    except OSError:
        return None


class DocumentRenderer:
    def __init__(self, *, font_path=None, docx_renderer='auto'):
        self.font_path = Path(font_path) if font_path is not None else Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts' / 'simsun.ttc'
        self.docx_renderer = _native_docx_renderer() if docx_renderer == 'auto' else docx_renderer

    def capabilities(self):
        return {'markdown': True, 'pdf': self.font_path.is_file(), 'docx_package': True,
            'docx_visual': self.docx_renderer is not None, 'font_path': str(self.font_path),
            'font_sha256': digest(self.font_path.read_bytes()) if self.font_path.is_file() else None}

    def _font(self, model):
        from fontTools.ttLib import TTFont
        if not self.font_path.is_file(): raise ValueError('FONT_UNAVAILABLE')
        with TTFont(str(self.font_path), fontNumber=0) as font:
            cmap = font.getBestCmap()
            text = ''.join(_block_text(b) for b in model.blocks)
            missing = sorted({ord(c) for c in text if not c.isspace() and ord(c) not in cmap})
            if missing: raise ValueError('FONT_GLYPH_MISSING:' + ','.join(hex(c) for c in missing))
            name = font['name'].getDebugName(1)
        return {'font_name': name, 'font_sha256': digest(self.font_path.read_bytes()), 'missing_glyphs': []}

    def render(self, model, format, *, evidence_dir):
        model = DocumentModel.model_validate(model.model_dump())
        if not model.complete: raise ValueError('DOCUMENT_PARSE_INCOMPLETE')
        if format not in ('markdown', 'docx', 'pdf'): raise ValueError('RENDER_FORMAT_UNSUPPORTED')
        root = Path(evidence_dir).resolve(); root.mkdir(parents=True, exist_ok=True)
        if root.is_symlink() or any(p.is_symlink() for p in root.parents): raise ValueError('RENDER_EVIDENCE_SCOPE_INVALID')
        if model.resources or any(b.resource_ids for b in model.blocks): raise ValueError('RENDER_RESOURCES_NOT_VERIFIED')
        if any(b.kind not in ('heading', 'paragraph', 'list_item', 'table_cell', 'code', 'separator') for b in model.blocks):
            raise ValueError('RENDER_BLOCK_KIND_UNSUPPORTED')
        if format == 'markdown': return self._markdown(model)
        font = self._font(model)
        if format == 'pdf': return self._pdf(model, root, font)
        return self._docx(model, root, font)

    def _markdown(self, model):
        from .parse import parse_document
        parts, index = [], 0
        while index < len(model.blocks):
            b = model.blocks[index]
            if b.kind == 'table_cell':
                number = b.metadata['table']; rows = {}
                while index < len(model.blocks) and model.blocks[index].kind == 'table_cell' and model.blocks[index].metadata['table'] == number:
                    cell = model.blocks[index]; rows.setdefault(cell.metadata['row'], {})[cell.metadata['col']] = cell.content.replace('|', '\\|')
                    index += 1
                table = []
                for rownum, row in sorted(rows.items()):
                    table.append('|' + '|'.join(row[c] for c in sorted(row)) + '|')
                    if rownum == min(rows): table.append('|' + '|'.join('---' for _ in row) + '|')
                parts.append('\n'.join(table) + '\n\n'); continue
            if b.kind == 'heading': parts.append('#'*b.metadata.get('level', 1) + ' ' + b.content + '\n\n')
            elif b.kind == 'code':
                if '```' in b.content: raise ValueError('MARKDOWN_FENCE_CONTENT_UNSUPPORTED')
                parts.append('```'+b.metadata.get('info', '')+'\n'+b.content.rstrip('\n')+'\n```\n\n')
            elif b.kind == 'list_item': parts.append('- ' + b.content + '\n\n')
            elif b.kind == 'separator': parts.append('---\n\n')
            else: parts.append(b.content + '\n\n')
            index += 1
        data = ''.join(parts).encode('utf-8')
        parsed = parse_document(data, 'markdown', source_version_id=model.source_version_id, language=model.language)
        if not parsed.complete or [(b.kind,b.content,b.metadata) for b in parsed.blocks] != [(b.kind,b.content,b.metadata) for b in model.blocks]:
            raise ValueError('MARKDOWN_CONTENT_ROUNDTRIP_FAILED')
        return RenderedDocument('markdown', model.language, data, {'status':'pass','code':'MARKDOWN_AST_ROUNDTRIP',
            'content_sha256':model.content_sha256,'byte_sha256':digest(data),'blocks':len(model.blocks),
            'visual_review':'not_required_text_artifact'})

    def _pdf(self, model, root, font):
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Preformatted, Spacer, Table, TableStyle
        from xml.sax.saxutils import escape
        fontname = 'Workbench-'+font['font_sha256'][:16]
        if fontname not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(fontname, str(self.font_path), subfontIndex=0))
        style = ParagraphStyle('body', fontName=fontname, fontSize=11, leading=17,
            spaceAfter=8, wordWrap='CJK', allowWidows=0, allowOrphans=0)
        output = io.BytesIO(); story = []; index = 0
        while index < len(model.blocks):
            b = model.blocks[index]
            if b.kind == 'table_cell':
                rows = {}; number = b.metadata['table']
                while index < len(model.blocks) and model.blocks[index].kind == 'table_cell' and model.blocks[index].metadata['table'] == number:
                    cell = model.blocks[index]
                    if len(cell.content) > 500: raise ValueError('TABLE_CELL_TOO_WIDE')
                    rows.setdefault(cell.metadata['row'], {})[cell.metadata['col']] = Paragraph(_pdf_rich(cell.content), style)
                    index += 1
                columns = len(next(iter(rows.values())))
                if columns > 8 or any(len(row) != columns for row in rows.values()): raise ValueError('TABLE_STRUCTURE_OR_WIDTH_UNSUPPORTED')
                table = Table([[row[c] for c in sorted(row)] for _, row in sorted(rows.items())], colWidths=[480/columns]*columns, repeatRows=1, hAlign='LEFT')
                table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),0.4,colors.grey), ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#f1f3f5')), ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6)]))
                story.extend([table, Spacer(1, 10)]); continue
            if b.kind == 'separator': story.append(Spacer(1,12))
            elif b.kind == 'code':
                if any(pdfmetrics.stringWidth(line, fontname, 9) > 480 for line in b.content.splitlines()): raise ValueError('CODE_LINE_TOO_WIDE')
                code = ParagraphStyle('code', parent=style, fontSize=9, leading=13)
                story.append(Preformatted(b.content, code))
            else:
                current = style
                if b.kind == 'heading': current = ParagraphStyle('heading', parent=style, fontSize=max(12,19-b.metadata.get('level',1)*2), leading=23, spaceBefore=10, keepWithNext=True)
                if b.kind=='list_item':
                    current=ParagraphStyle('bullet',parent=current,leftIndent=14,bulletIndent=0)
                    story.append(Paragraph(_pdf_rich(b.content),current,bulletText='-'))
                else:
                    story.append(Paragraph(_pdf_rich(b.content), current))
            index += 1
        SimpleDocTemplate(output, pagesize=(595.276,841.89), rightMargin=54,leftMargin=54,topMargin=54,bottomMargin=54,
            title='Report',author='').build(story)
        data = output.getvalue()
        qa, images = verify_pdf(data, model, root)
        qa.update(font)
        return RenderedDocument('pdf', model.language, data, qa, images)

    def _docx(self, model, root, font):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt, Cm
        from docx.opc.constants import RELATIONSHIP_TYPE
        from .parse import parse_document
        document = Document(); section = document.sections[0]
        section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Cm(1.9)
        normal = document.styles['Normal']; normal.font.name=font['font_name']; normal.font.size=Pt(11)
        normal.element.rPr.rFonts.set(qn('w:eastAsia'), font['font_name'])
        for heading in range(1,7):
            heading_style=document.styles['Heading '+str(heading)]
            heading_style.font.name=font['font_name']
            heading_style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font['font_name'])
        def add_content(paragraph,text):
            offset=0
            for match in re.finditer(r'\[([^]]+)\]\((https?://[^)\s]+|mailto:[^)\s]+)\)',text):
                paragraph.add_run(_visible(text[offset:match.start()]))
                relation=paragraph.part.relate_to(match[2],RELATIONSHIP_TYPE.HYPERLINK,is_external=True)
                hyperlink=OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'),relation)
                run=OxmlElement('w:r'); value=OxmlElement('w:t'); value.text=match[1]; run.append(value); hyperlink.append(run)
                paragraph._p.append(hyperlink); offset=match.end()
            paragraph.add_run(_visible(text[offset:]))
        index = 0
        while index < len(model.blocks):
            b = model.blocks[index]
            if b.kind == 'table_cell':
                rows={}; number=b.metadata['table']
                while index<len(model.blocks) and model.blocks[index].kind=='table_cell' and model.blocks[index].metadata['table']==number:
                    cell=model.blocks[index]; rows.setdefault(cell.metadata['row'],{})[cell.metadata['col']]=cell.content; index+=1
                columns=len(next(iter(rows.values())))
                if any(len(r)!=columns for r in rows.values()): raise ValueError('TABLE_STRUCTURE_UNSUPPORTED')
                table=document.add_table(rows=len(rows),cols=columns); table.style='Table Grid'
                for ri, (_,row) in enumerate(sorted(rows.items())):
                    for ci,text in sorted(row.items()): add_content(table.cell(ri,ci).paragraphs[0],text)
                header=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(header)
                continue
            if b.kind=='heading': add_content(document.add_heading('',level=b.metadata.get('level',1)),b.content)
            elif b.kind=='separator': document.add_paragraph('')
            elif b.kind=='code': document.add_paragraph(b.content)
            else: add_content(document.add_paragraph('',style='List Bullet' if b.kind=='list_item' else None),b.content)
            index+=1
        stream=io.BytesIO(); document.save(stream); data=stream.getvalue()
        parsed=parse_document(data,'docx',source_version_id=model.source_version_id,language=model.language)
        expected=[_compact(_block_text(b)) for b in model.blocks if b.kind!='separator']
        actual=[_compact(b.content) for b in parsed.blocks]
        if not parsed.complete or expected!=actual: raise ValueError('DOCX_CONTENT_ROUNDTRIP_FAILED')
        qa={'status':'blocked','code':'DOCX_VISUAL_RENDERER_UNAVAILABLE','byte_sha256':digest(data),
            'content_sha256':model.content_sha256,'package_verified':True,'visual_review':'required',**font}
        if self.docx_renderer is None: return RenderedDocument('docx',model.language,data,qa)
        # A configured native renderer must return actual converted page bytes.
        preview=self.docx_renderer(data, root)
        report,images=verify_pdf(preview,model,root)
        report.update(font,byte_sha256=digest(data),preview_sha256=digest(preview),package_verified=True)
        return RenderedDocument('docx',model.language,data,report,images,preview)


def verify_pdf(data, model, evidence_dir):
    import pymupdf
    images=[]; text=[]
    with pymupdf.open(stream=data,filetype='pdf') as pdf:
        if not len(pdf): raise ValueError('PDF_EMPTY')
        for index,page in enumerate(pdf):
            value=page.get_text(); text.append(value)
            if not value.strip(): raise ValueError('PDF_BLANK_PAGE')
            for item in page.get_text('dict')['blocks']:
                if item.get('type')!=0: continue
                for line in item['lines']:
                    for span in line['spans']:
                        x0,y0,x1,y1=span['bbox']
                        if x0<0 or y0<0 or x1>page.rect.width+0.1 or y1>page.rect.height+0.1:
                            raise ValueError('PDF_TEXT_OUTSIDE_PAGE')
                        if '\ufffd' in span['text'] or '\u25a1' in span['text']: raise ValueError('PDF_MISSING_GLYPH')
            image=page.get_pixmap(matrix=pymupdf.Matrix(1.25,1.25))
            path=Path(evidence_dir)/f'{model.language}-{digest(data)[:12]}-page-{index+1}.png'
            image.save(str(path)); images.append(path)
        observed=_compact('\n'.join(text))
        actual_links=[link.get('uri') for page in pdf for link in page.get_links() if link.get('uri')]
        required_links=[uri for b in model.blocks for uri in b.metadata.get('links',[]) if uri.startswith(('http://','https://','mailto:'))]
        if any(uri not in actual_links for uri in required_links): raise ValueError('PDF_REFERENCE_LINK_MISSING')
        cursor=0; block_positions=[]
        for b in model.blocks:
            expected=_compact(_block_text(b))
            if not expected: continue
            position=observed.find(expected,cursor)
            if position<0: raise ValueError('PDF_CONTENT_MISSING:'+b.block_id)
            cursor=position+len(expected)
            block_positions.append({'block_id':b.block_id,'text_start':position,'text_end':cursor})
        qa={'status':'pass','code':'PDF_CONTENT_FONT_GEOMETRY_VERIFIED','byte_sha256':digest(data),
            'content_sha256':model.content_sha256,'pages':len(pdf),'page_images_sha256':[digest(p.read_bytes()) for p in images],
            'visual_review':'required','text_verified':True,'geometry_verified':True,'ordered_block_positions':block_positions}
    return qa,tuple(images)
