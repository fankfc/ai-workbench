import io
import zipfile

from docx import Document
import pytest


def test_markdown_ast_preserves_code_table_link_and_appendix_without_silent_html():
    from app.documents.parse import parse_document
    text = '# Report\n\n[Source](https://example.com)\n\n|Month|Amount|\n|---|---|\n|Jan|30|\n\n```python\nprint(30)\n```\n\n## Appendix\n\nFull appendix.\n'
    model = parse_document(text.encode(), 'markdown', source_version_id='s1', language='en')
    assert model.complete
    assert [b.kind for b in model.blocks] == ['heading', 'paragraph', 'table_cell', 'table_cell', 'table_cell', 'table_cell', 'code', 'heading', 'paragraph']
    assert model.blocks[6].content == 'print(30)\n'
    assert not model.blocks[6].translatable
    assert any(b.metadata.get('links') == ['https://example.com'] for b in model.blocks)
    assert all(b.source_locator and b.source_sha256 for b in model.blocks)
    html = parse_document(b'# Title\n\n<script>bad()</script>', 'markdown', source_version_id='s1', language='en')
    assert not html.complete
    assert 'UNSUPPORTED_HTML' in [i.code for i in html.issues]


def test_docx_extracts_body_table_and_footnote_instead_of_ignoring_package_parts():
    from app.documents.parse import parse_document
    doc = Document()
    doc.add_heading('Report', 1)
    doc.add_paragraph('Body 30.')
    doc.add_table(rows=1, cols=2).cell(0, 1).text = '40'
    stream = io.BytesIO(); doc.save(stream)
    source = io.BytesIO()
    with zipfile.ZipFile(source, 'w') as out, zipfile.ZipFile(io.BytesIO(stream.getvalue())) as original:
        for name in original.namelist():
            out.writestr(name, original.read(name))
        out.writestr('word/footnotes.xml', '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>Required footnote 10.</w:t></w:r></w:p></w:footnote></w:footnotes>')
    model = parse_document(source.getvalue(), 'docx', source_version_id='s', language='en')
    assert model.complete
    assert any(b.kind == 'footnote' and b.content == 'Required footnote 10.' for b in model.blocks)
    assert len([b for b in model.blocks if b.kind == 'table_cell']) == 2


def test_scanned_pdf_and_multiple_columns_require_explicit_input():
    import pymupdf
    from app.documents.parse import parse_document
    scan = pymupdf.open(); scan.new_page()
    model = parse_document(scan.tobytes(), 'pdf', source_version_id='s', language='en')
    assert not model.complete
    assert 'PDF_OCR_REQUIRED' in [i.code for i in model.issues]
    columns = pymupdf.open(); page = columns.new_page()
    page.insert_text((40, 60), 'Left column body')
    page.insert_text((330, 60), 'Right column body')
    model = parse_document(columns.tobytes(), 'pdf', source_version_id='s', language='en')
    assert not model.complete
    assert 'PDF_READING_ORDER_AMBIGUOUS' in [i.code for i in model.issues]


@pytest.mark.parametrize('fmt,data', [('doc', b'legacy'), ('docx', b'broken'), ('pdf', b'broken')])
def test_unsupported_and_corrupt_inputs_never_look_complete(fmt, data):
    from app.documents.parse import parse_document
    model = parse_document(data, fmt, source_version_id='s', language='en')
    assert not model.complete
    assert model.issues


@pytest.mark.parametrize('text', [b'1. First action.\n2. Second action.\n',
    b'- First action.\n  - Required child step.\n',b'- First action.\n\n  Required continuation.\n'])
def test_unsupported_list_structure_is_explicit_instead_of_flattened(text):
    from app.documents.parse import parse_document
    model=parse_document(text,'markdown',source_version_id='s',language='en')
    assert not model.complete
    assert any(i.code.startswith('MARKDOWN_LIST_') for i in model.issues)


def test_simple_flat_bullets_remain_supported_with_structure_metadata():
    from app.documents.parse import parse_document
    model=parse_document(b'- First action.\n- Second action.\n','markdown',source_version_id='s',language='en')
    assert model.complete
    assert all(b.metadata.get('list_type')=='unordered' and b.metadata.get('list_depth')==1 for b in model.blocks)
